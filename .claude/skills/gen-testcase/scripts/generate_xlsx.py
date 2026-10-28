#!/usr/bin/env python3
"""Generate a formatted QA test-case .xlsx from the gen-testcase JSON contract.

Input JSON shape:
{
  "sheetName": "[TEPO] File Upload",
  "appPrefix": "TEPO",
  "testCases": [ { "Test Case ID": "...", "Module": "...", ... }, ... ]
}

Usage:
  python3 generate_xlsx.py --json testcases_XXX.json --output outputs/testcases/test-case-xxx.xlsx
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNS = [
    "Test Case ID",
    "Module",
    "Test Scenario",
    "Priority",
    "Preconditions",
    "Steps",
    "Data",
    "Expected Result",
    "Status",
    "Actual Result",
    "Notes",
]

TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
DATA_FILL_A = PatternFill("solid", fgColor="FFFFFF")
DATA_FILL_B = PatternFill("solid", fgColor="DEEAF1")
PRIORITY_FILL = PatternFill("solid", fgColor="FCE4D6")

STATUS_FILLS = {
    "Pending": PatternFill("solid", fgColor="F7ECE6"),
    "Pass": PatternFill("solid", fgColor="D4EDBC"),
    "Fail": PatternFill("solid", fgColor="FFCFC9"),
    "Blocked": PatternFill("solid", fgColor="FFE5A0"),
    "N/A": PatternFill("solid", fgColor="E6E6E6"),
}

THIN_GRAY = Side(style="thin", color="D9D9D9")
DATA_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def build_workbook(sheet_name: str, app_prefix: str, test_cases: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    n_cols = len(COLUMNS)

    # Row 1: title row, merged across all columns
    title_text = sheet_name
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Row 2: header row
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=col_name)
        cell.font = Font(name="Arial", size=11, bold=True, color="000000")
        cell.fill = HEADER_FILL
        cell.border = DATA_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data rows
    for row_offset, tc in enumerate(test_cases):
        row = 3 + row_offset
        base_fill = DATA_FILL_A if row_offset % 2 == 0 else DATA_FILL_B
        status_value = tc.get("Status", "Pending") or "Pending"
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            value = tc.get(col_name, "")
            cell = ws.cell(row=row, column=col_idx, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.border = DATA_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if col_name == "Status":
                cell.fill = STATUS_FILLS.get(status_value, STATUS_FILLS["Pending"])
            elif col_name == "Priority":
                cell.fill = PRIORITY_FILL
            else:
                cell.fill = base_fill

    last_row = 2 + len(test_cases)

    # Dropdown validation: Status
    if last_row >= 3:
        status_col = COLUMNS.index("Status") + 1
        status_letter = get_column_letter(status_col)
        dv_status = DataValidation(
            type="list",
            formula1='"Pending,Pass,Fail,Blocked,N/A"',
            allow_blank=True,
        )
        ws.add_data_validation(dv_status)
        dv_status.add(f"{status_letter}3:{status_letter}{last_row}")

        priority_col = COLUMNS.index("Priority") + 1
        priority_letter = get_column_letter(priority_col)
        dv_priority = DataValidation(
            type="list",
            formula1='"Critical,High,Medium,Low"',
            allow_blank=True,
        )
        ws.add_data_validation(dv_priority)
        dv_priority.add(f"{priority_letter}3:{priority_letter}{last_row}")

    # Column widths: auto-fit-ish, minimum 15
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(col_name)
        for row_offset in range(len(test_cases)):
            value = str(test_cases[row_offset].get(col_name, ""))
            longest_line = max((len(line) for line in value.split("\n")), default=0)
            max_len = max(max_len, longest_line)
        ws.column_dimensions[letter].width = max(15, min(max_len + 2, 60))

    # Freeze title + header rows
    ws.freeze_panes = "A3"

    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--json", required=True, help="Path to testcases_*.json")
    parser.add_argument("--output", required=True, help="Path to write the .xlsx file")
    args = parser.parse_args()

    with open(args.json, encoding="utf-8") as f:
        payload = json.load(f)

    for required_field in ("sheetName", "appPrefix", "testCases"):
        if required_field not in payload:
            print(
                f"Thieu truong: JSON can co: sheetName, appPrefix, testCases "
                f"(thieu '{required_field}')",
                file=sys.stderr,
            )
            return 1

    wb = build_workbook(payload["sheetName"], payload["appPrefix"], payload["testCases"])
    wb.save(args.output)
    print(f"Da tao: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
